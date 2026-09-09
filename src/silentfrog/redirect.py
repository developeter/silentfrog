"""Massive Redirect Check — bulk verification of an Old URL -> New URL map.

Reads a spreadsheet of migration redirects and, for every row, walks the
redirect chain hop by hop instead of letting ``requests`` swallow it. Walking
manually is what makes the audit useful: each hop's status code is recorded, a
loop is reported as a loop (not as a generic error), the hop budget is
explicit, and — like every other fetch in the app — each hop is vetted against
the SSRF guard before it is dialled.
"""

from __future__ import annotations

import logging
import re
import socket
import threading
import time
from collections.abc import Callable, Iterator
from concurrent.futures import Future, ThreadPoolExecutor, as_completed
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import unquote, urljoin, urlsplit

import pandas as pd
import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry

from .crawl_options import DEFAULT_USER_AGENT
from .robots_simulator import parse_robots
from .ssrf import is_public_address

__all__ = [
    "RedirectCheckOptions",
    "RedirectOutcome",
    "RedirectRunResult",
    "RedirectSummary",
    "RedirectThrottle",
    "check_redirects",
]

USER_AGENT = DEFAULT_USER_AGENT
_logger = logging.getLogger(__name__)

OLD_URL_COLUMN = "Old URL"
NEW_URL_COLUMN = "New URL"
_INPUT_COLUMNS = (OLD_URL_COLUMN, NEW_URL_COLUMN)
# Header spellings a real migration sheet uses, matched after casefolding and
# dropping separators, so "old_url", "Old Url" and "OLD-URL" all land here.
_OLD_HEADER_ALIASES = frozenset({"oldurl", "old", "from", "fromurl", "source", "sourceurl", "urlold"})
_NEW_HEADER_ALIASES = frozenset({"newurl", "new", "to", "tourl", "target", "targeturl", "urlnew", "destination"})

STATUS_COLUMN = "Status Code"
ISSUE_COLUMN = "Issue"
_RESULT_COLUMNS = (
    STATUS_COLUMN,
    "Final URL",
    "Redirect Correct",
    "Redirect Chain Length",
    "Redirect Chain URLs",
    "Redirect Chain Statuses",
    ISSUE_COLUMN,
)

# Issue vocabulary. Empty means "nothing to report", so the column stays quiet
# on the rows a user does not need to look at.
ISSUE_NONE = ""
ISSUE_WRONG_TARGET = "Wrong target"
ISSUE_CHAIN = "Redirect chain"
ISSUE_LOOP = "Redirect loop"
ISSUE_TOO_MANY_HOPS = "Too many hops"
ISSUE_ROBOTS = "Blocked by robots.txt"
ISSUE_PRIVATE = "Blocked: private address"
ISSUE_INVALID = "Invalid URL"
ISSUE_NO_TARGET = "No New URL"
ISSUE_SSL = "SSL error"
ISSUE_ERROR = "Request error"
ISSUE_NOT_CHECKED = "Not checked"

# Escapes that must survive normalisation: an encoded delimiter is data, not
# a delimiter. / ? # % & = in that order.
_RESERVED_ESCAPE = re.compile(r"(%2[Ff]|%3[Ff]|%23|%25|%26|%3[Dd])")

_DEFAULT_PORTS = {"http": 80, "https": 443}
_MAX_LOG_URL = 120
_MAX_OUTPUT_ATTEMPTS = 20
# Worker-pool ceiling. The pool is built once at this size and real
# concurrency is gated by RedirectThrottle, so Threads stays adjustable
# mid-run; it matches the GUI spinbox maximum.
MAX_PARALLEL = 20


@dataclass(frozen=True, slots=True)
class RedirectCheckOptions:
    """One typed bundle instead of a bool-flag pile across the call chain."""

    timeout: int = 10
    max_workers: int = 5
    max_hops: int = 10
    respect_robots: bool = True
    verify_ssl: bool = True
    # Mirrors the crawler's ``CrawlOptions.allow_private_network`` opt-in: off
    # by default, so a redirect sheet cannot aim the tool at loopback or an
    # intranet host the operator never meant to touch.
    allow_private_network: bool = False


@dataclass(frozen=True, slots=True)
class RedirectHop:
    url: str
    status: int


@dataclass(frozen=True, slots=True)
class RedirectOutcome:
    """Everything one row's chain walk produced, before it is flattened into
    spreadsheet cells."""

    status: int | str
    final_url: str
    hops: tuple[RedirectHop, ...]
    issue: str

    @property
    def chain_length(self) -> int:
        """Redirect hops, i.e. responses before the final one."""
        return max(0, len(self.hops) - 1)

    @property
    def chain_urls(self) -> str:
        return " -> ".join(hop.url for hop in self.hops)

    @property
    def chain_statuses(self) -> str:
        return " -> ".join(str(hop.status) for hop in self.hops)


@dataclass(frozen=True, slots=True)
class RedirectSummary:
    """Headline counts, so the GUI can say what happened instead of "done"."""

    total: int = 0
    correct: int = 0
    wrong: int = 0
    chains: int = 0
    loops: int = 0
    errors: int = 0


@dataclass(frozen=True, slots=True)
class RedirectRunResult:
    output_path: Path
    summary: RedirectSummary
    cancelled: bool


class RedirectThrottle:
    """Live speed control: how many rows run at once, and how long to wait
    before each request.

    Both are adjustable WHILE the run is in flight. A bulk redirect check only
    finds out it is too fast once the target starts refusing it, and by then
    restarting a several-thousand-row audit to change a number is the worst
    possible answer — so the knobs stay live instead of being frozen at Start.
    The worker pool is sized once to ``max_parallel``; concurrency below that
    is enforced here rather than by resizing the pool, which is not something
    ThreadPoolExecutor supports.
    """

    def __init__(self, parallel: int, delay_ms: int = 0, max_parallel: int = MAX_PARALLEL) -> None:
        self._max_parallel = max(1, max_parallel)
        self._parallel = self._clamped(parallel)
        self._delay = max(0, delay_ms) / 1000
        self._active = 0
        self._condition = threading.Condition()

    def _clamped(self, value: int) -> int:
        return max(1, min(int(value), self._max_parallel))

    @property
    def max_parallel(self) -> int:
        return self._max_parallel

    def set_parallel(self, value: int) -> None:
        with self._condition:
            self._parallel = self._clamped(value)
            # Raising the limit must wake threads already queued, or the new
            # value would only take effect as running rows happen to finish.
            self._condition.notify_all()

    def set_delay_ms(self, value: int) -> None:
        with self._condition:
            self._delay = max(0, int(value)) / 1000

    def pause_before_request(self) -> None:
        with self._condition:
            delay = self._delay
        if delay:
            time.sleep(delay)

    @contextmanager
    def slot(self) -> Iterator[None]:
        """Hold one of the live concurrency slots for the duration of a row."""
        with self._condition:
            while self._active >= self._parallel:
                self._condition.wait(0.2)
            self._active += 1
        try:
            yield
        finally:
            with self._condition:
                self._active -= 1
                self._condition.notify()


# --------------------------------------------------------------------------- #
# URL comparison
# --------------------------------------------------------------------------- #
def _decoded(text: str) -> str:
    """Percent-decode everything EXCEPT the delimiters, which carry meaning.

    ``/caff%C3%A8`` and ``/caffè`` are the same resource; ``/a%2Fb`` and
    ``/a/b`` are not — an encoded slash is a path segment containing a slash,
    not a segment boundary. Decoding blindly would silently pass a redirect
    that landed on the wrong resource. Splitting on the reserved escapes first
    keeps them verbatim while each surrounding run still decodes as a whole,
    which multi-byte UTF-8 needs (``%C3%A8`` is one character, not two).
    """
    parts = _RESERVED_ESCAPE.split(text)
    return "".join(part.upper() if index % 2 else unquote(part) for index, part in enumerate(parts))


def _comparable_url(url: str) -> str:
    """A form two URLs can be compared in without cosmetic false negatives.

    ``requests`` reports the percent-encoded final URL while the sheet holds
    the human spelling, so ``/caff%C3%A8`` and ``/caffè`` must compare equal;
    host case and a default port are equally meaningless. Path case is NOT
    folded — it is significant on most origins.
    """
    parsed = urlsplit(url.strip())
    scheme = parsed.scheme.lower()
    host = (parsed.hostname or "").lower()
    port = parsed.port
    suffix = f":{port}" if port and port != _DEFAULT_PORTS.get(scheme) else ""
    path = _decoded(parsed.path).rstrip("/")
    return f"{scheme}://{host}{suffix}{path}?{_decoded(parsed.query)}"


def _is_http_url(url: str) -> bool:
    parsed = urlsplit(url.strip())
    return parsed.scheme in _DEFAULT_PORTS and bool(parsed.hostname)


def _cell_url(value: object) -> str:
    """A spreadsheet cell as a URL string. Blanks and pandas' NaN both become
    "", so a missing cell never reaches the network as the literal "nan"."""
    if value is None or (isinstance(value, float) and value != value):
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none"} else text


# --------------------------------------------------------------------------- #
# SSRF vetting (synchronous path)
# --------------------------------------------------------------------------- #
def _resolves_public(url: str) -> bool:
    """True when every address the host resolves to is globally routable.

    The async crawler pins the connection to the vetted IP; ``requests``
    re-resolves on its own, so this pre-flight leaves a narrow DNS-rebinding
    window. It still stops the realistic case — a sheet row aimed at loopback,
    an intranet host, or the cloud-metadata endpoint. A resolution failure is
    not a verdict: let the request run and report the real transport error.
    """
    parsed = urlsplit(url)
    host = parsed.hostname
    if not host:
        return False
    port = parsed.port or _DEFAULT_PORTS.get(parsed.scheme, 80)
    try:
        resolved = socket.getaddrinfo(host, port, proto=socket.IPPROTO_TCP)
    except OSError:
        return True
    return all(is_public_address(info[4][0]) for info in resolved)


# --------------------------------------------------------------------------- #
# robots.txt, cached per origin
# --------------------------------------------------------------------------- #
class _RobotsCache:
    """One robots.txt fetch per origin per run.

    The previous implementation re-fetched robots.txt for every row, so a
    20k-row single-host migration sent 20k redundant requests and doubled the
    traffic the audit was meant to keep polite.
    """

    def __init__(self, timeout: int, verify_ssl: bool) -> None:
        self._timeout = timeout
        self._verify_ssl = verify_ssl
        self._bodies: dict[str, str] = {}
        self._lock = threading.Lock()

    def allows(self, session: requests.Session, url: str) -> bool:
        parsed = urlsplit(url)
        return parse_robots(self._body(session, f"{parsed.scheme}://{parsed.netloc}")).allows("*", url)

    def _body(self, session: requests.Session, origin: str) -> str:
        with self._lock:
            cached = self._bodies.get(origin)
            if cached is None:
                cached = self._fetch(session, origin)
                self._bodies[origin] = cached
            return cached

    def _fetch(self, session: requests.Session, origin: str) -> str:
        # Fail open: an unreachable or erroring robots.txt must not turn every
        # row of the sheet into a false "blocked".
        try:
            response = session.get(f"{origin}/robots.txt", timeout=self._timeout, verify=self._verify_ssl)
        except requests.RequestException:
            return ""
        return response.text if response.status_code < 400 else ""


# --------------------------------------------------------------------------- #
# HTTP session, one per worker thread
# --------------------------------------------------------------------------- #
def _new_session(verify_ssl: bool) -> requests.Session:
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    session.verify = verify_ssl
    retry = Retry(
        total=3,
        backoff_factor=0.3,
        status_forcelist=(500, 502, 503, 504),
        allowed_methods=("GET", "HEAD"),
        raise_on_status=False,
    )
    adapter = HTTPAdapter(max_retries=retry, pool_maxsize=20)
    session.mount("https://", adapter)
    session.mount("http://", adapter)
    return session


class _SessionPool:
    """A session per worker thread, not per row: the connection pool and the
    retry policy only pay off when a session outlives a single request."""

    def __init__(self, verify_ssl: bool) -> None:
        self._verify_ssl = verify_ssl
        self._local = threading.local()

    def session(self) -> requests.Session:
        existing = getattr(self._local, "session", None)
        if existing is not None:
            return existing
        created = _new_session(self._verify_ssl)
        self._local.session = created
        return created


# --------------------------------------------------------------------------- #
# Chain walking
# --------------------------------------------------------------------------- #
def _blocked_reason(
    session: requests.Session,
    url: str,
    options: RedirectCheckOptions,
    robots: _RobotsCache,
) -> str:
    if not _is_http_url(url):
        return ISSUE_INVALID
    if not options.allow_private_network and not _resolves_public(url):
        return ISSUE_PRIVATE
    if options.respect_robots and not robots.allows(session, url):
        return ISSUE_ROBOTS
    return ISSUE_NONE


def _outcome(hops: list[RedirectHop], final_url: str, issue: str) -> RedirectOutcome:
    status: int | str = hops[-1].status if hops else ""
    return RedirectOutcome(status=status, final_url=final_url, hops=tuple(hops), issue=issue)


def _next_hop(response: requests.Response, current_url: str) -> str:
    location = response.headers.get("Location", "")
    if not location or not response.is_redirect:
        return ""
    return urljoin(current_url, location.strip())


def _request(
    session: requests.Session,
    url: str,
    options: RedirectCheckOptions,
) -> tuple[requests.Response | None, str]:
    try:
        # The body is read (no ``stream=True``) purely to keep the connection
        # reusable: urllib3 can only return a socket to the pool once the
        # response is fully consumed, and closing an unread streamed response
        # discards it. Streaming here to skip the landing page cost one fresh
        # TCP+TLS handshake PER ROW, which on a several-thousand-row
        # single-host migration reads as an attack and gets the run dropped at
        # the edge (ConnectTimeout). Bandwidth is the cheaper thing to spend.
        response = session.get(
            url,
            allow_redirects=False,
            timeout=options.timeout,
            verify=options.verify_ssl,
        )
    except requests.exceptions.SSLError:
        return None, ISSUE_SSL
    except requests.RequestException as exc:
        return None, f"{ISSUE_ERROR}: {exc.__class__.__name__}"
    return response, ISSUE_NONE


def _follow_chain(
    session: requests.Session,
    start_url: str,
    options: RedirectCheckOptions,
    robots: _RobotsCache,
    throttle: RedirectThrottle | None = None,
) -> RedirectOutcome:
    hops: list[RedirectHop] = []
    seen: set[str] = set()
    url = start_url
    while len(hops) <= options.max_hops:
        blocked = _blocked_reason(session, url, options, robots)
        if blocked:
            return _outcome(hops, url, blocked)
        if throttle is not None:
            throttle.pause_before_request()
        response, failure = _request(session, url, options)
        if response is None:
            return _outcome(hops, url, failure)
        hops.append(RedirectHop(url, response.status_code))
        seen.add(_comparable_url(url))
        url = _next_hop(response, url)
        if not url:
            return _outcome(hops, hops[-1].url, ISSUE_NONE)
        if _comparable_url(url) in seen:
            return _outcome(hops, url, ISSUE_LOOP)
    return _outcome(hops, url, ISSUE_TOO_MANY_HOPS)


def _verdict(outcome: RedirectOutcome, expected_url: str) -> tuple[bool, str]:
    """(redirect correct, issue). A transport-level issue already decided the
    row, so only a clean walk is judged against the expected target."""
    if outcome.issue:
        return False, outcome.issue
    if not expected_url:
        return False, ISSUE_NO_TARGET
    if _comparable_url(outcome.final_url) != _comparable_url(expected_url):
        return False, ISSUE_WRONG_TARGET
    return True, ISSUE_CHAIN if outcome.chain_length > 1 else ISSUE_NONE


# --------------------------------------------------------------------------- #
# Row plumbing
# --------------------------------------------------------------------------- #
@dataclass(frozen=True, slots=True)
class _RedirectJob:
    index: int
    old_url: str
    new_url: str


@dataclass(frozen=True, slots=True)
class _RedirectRowResult:
    index: int
    old_url: str
    status: int | str
    final_url: str
    is_correct: bool
    chain_length: int
    chain_urls: str
    chain_statuses: str
    issue: str


def _row_result(job: _RedirectJob, outcome: RedirectOutcome, verdict: tuple[bool, str]) -> _RedirectRowResult:
    is_correct, issue = verdict
    return _RedirectRowResult(
        index=job.index,
        old_url=job.old_url,
        status=outcome.status,
        final_url=outcome.final_url,
        is_correct=is_correct,
        chain_length=outcome.chain_length,
        chain_urls=outcome.chain_urls,
        chain_statuses=outcome.chain_statuses,
        issue=issue,
    )


def _apply_redirect_result(df: pd.DataFrame, result: _RedirectRowResult) -> None:
    # Positional assignment: a sheet whose index is not a clean RangeIndex (or
    # holds duplicate labels) would otherwise splatter one row's result across
    # every row sharing that label.
    values = (
        result.status,
        result.final_url,
        "Yes" if result.is_correct else "No",
        result.chain_length,
        result.chain_urls,
        result.chain_statuses,
        result.issue,
    )
    for column, value in zip(_RESULT_COLUMNS, values, strict=True):
        df.iloc[result.index, df.columns.get_loc(column)] = value


# --------------------------------------------------------------------------- #
# Spreadsheet loading
# --------------------------------------------------------------------------- #
def _header_key(value: object) -> str:
    return "".join(ch for ch in str(value).lower() if ch.isalnum())


def _match_header(columns: list[object]) -> tuple[object, object] | None:
    """The (old, new) column labels, matched case- and separator-insensitively.
    ``None`` when the sheet has no recognisable header row at all.

    Matching exactly ONE of the two is the dangerous case: the sheet clearly
    HAS a header, so falling back to "first two columns" would read that header
    as data and could pick the wrong columns entirely. Better to stop and say
    which half was understood than to produce a confident wrong report.
    """
    old = next((col for col in columns if _header_key(col) in _OLD_HEADER_ALIASES), None)
    new = next((col for col in columns if _header_key(col) in _NEW_HEADER_ALIASES), None)
    if old is not None and new is not None:
        return old, new
    if old is None and new is None:
        return None
    if old is None:
        raise ValueError(f"Found a new-URL column but no old-URL column. Rename one to '{OLD_URL_COLUMN}'.")
    raise ValueError(f"Found an old-URL column but no new-URL column. Rename one to '{NEW_URL_COLUMN}'.")


def _read_any(path: Path, headerless: bool = False) -> pd.DataFrame:
    """CSV or Excel by extension. CSV matters because the Site Crawl redirect
    map (G13) exports exactly that, and this tool could not read its own
    sibling's output."""
    header = None if headerless else "infer"
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path, dtype=str, keep_default_na=False, header=header)
    return pd.read_excel(path, header=None if headerless else 0)


def _load_redirect_frame(path: Path) -> pd.DataFrame:
    """A two-column ``Old URL`` / ``New URL`` frame with a clean RangeIndex.

    Headered sheets are matched by alias; a headerless sheet falls back to its
    first two columns *by position*. The old code handed pandas two ``names``
    instead, which silently promoted the real Old URL column to the index and
    shifted every value one column left as soon as a third column (a note, a
    status) was present — the tool then checked the wrong URLs without a word.
    """
    probe = _read_any(path)
    matched = _match_header(list(probe.columns))
    frame = probe[[matched[0], matched[1]]] if matched else _read_any(path, headerless=True).iloc[:, :2]
    if frame.shape[1] < 2:
        raise ValueError("The file needs two columns: the old URL and the new URL.")
    frame = frame.copy()
    frame.columns = list(_INPUT_COLUMNS)
    return _clean_rows(frame)


def _clean_rows(frame: pd.DataFrame) -> pd.DataFrame:
    """Normalise the two URL cells and reset the index. Rows are KEPT even
    when blank — a migration sheet is audited row-for-row against its source,
    so a dropped row is worse than a row flagged "Invalid URL"."""
    cleaned = frame.assign(
        **{
            OLD_URL_COLUMN: [_cell_url(value) for value in frame[OLD_URL_COLUMN]],
            NEW_URL_COLUMN: [_cell_url(value) for value in frame[NEW_URL_COLUMN]],
        }
    ).reset_index(drop=True)
    if cleaned.empty:
        raise ValueError("No rows found. Expected an 'Old URL' and a 'New URL' column.")
    return cleaned


def _ensure_result_columns(df: pd.DataFrame) -> None:
    # Always reset: re-running on a previous report must not leave last run's
    # verdicts standing next to this run's.
    for column in _RESULT_COLUMNS:
        df[column] = ""


def _redirect_jobs(df: pd.DataFrame) -> list[_RedirectJob]:
    old_urls = list(df[OLD_URL_COLUMN])
    new_urls = list(df[NEW_URL_COLUMN])
    pairs = zip(old_urls, new_urls, strict=True)
    return [_RedirectJob(index, str(old), str(new)) for index, (old, new) in enumerate(pairs)]


# --------------------------------------------------------------------------- #
# Excel output
# --------------------------------------------------------------------------- #
def _column_index(worksheet, header: str) -> int | None:
    for index, cell in enumerate(worksheet[1], start=1):
        if cell.value == header:
            return index
    return None


def _row_needs_attention(row, status_idx: int | None, issue_idx: int | None) -> bool:
    """Any row a user must act on, not just a literal 404: a broken status or
    a reported issue. The old rule tinted 404s only, so 500s, loops and
    wrong-target rows looked exactly like a clean redirect."""
    status = str(row[status_idx - 1].value or "") if status_idx else ""
    issue = str(row[issue_idx - 1].value or "") if issue_idx else ""
    return bool(issue) or status[:1] in {"4", "5"} or not status


def _style_redirect_sheet(worksheet) -> None:
    from openpyxl.styles import Font, PatternFill

    default_font = Font(name="Arial", size=10)
    red_fill = PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid")
    status_idx = _column_index(worksheet, STATUS_COLUMN)
    issue_idx = _column_index(worksheet, ISSUE_COLUMN)
    for row in worksheet.iter_rows():
        _style_row(row, default_font, red_fill if _row_needs_attention(row, status_idx, issue_idx) else None)
    _finish_sheet(worksheet)


def _style_row(row, font, fill) -> None:
    for cell in row:
        cell.font = font
        if fill is not None:
            cell.fill = fill


def _finish_sheet(worksheet) -> None:
    """Freeze the header, add an autofilter and give the URL columns room —
    the report is read by scrolling and filtering, not by admiring it."""
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[1]:
        worksheet.column_dimensions[cell.column_letter].width = 52 if cell.column <= 2 else 18


def _write_redirect_results(df: pd.DataFrame, output_path: Path) -> Path:
    """Write the report, falling back to a numbered sibling when the target is
    locked (the classic "results file still open in Excel"). A finished run's
    results are never discarded because of the very last step."""
    fallbacks = (output_path.with_stem(f"{output_path.stem}_{n}") for n in range(1, _MAX_OUTPUT_ATTEMPTS))
    for candidate in (output_path, *fallbacks):
        try:
            return _write_workbook(df, candidate)
        except PermissionError:
            continue
    raise PermissionError(f"Could not write {output_path} — close it in Excel and run again.")


def _write_workbook(df: pd.DataFrame, path: Path) -> Path:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Results")
        _style_redirect_sheet(writer.sheets["Results"])
    return path


# --------------------------------------------------------------------------- #
# Summary
# --------------------------------------------------------------------------- #
_ERROR_ISSUES = frozenset({ISSUE_SSL, ISSUE_INVALID, ISSUE_PRIVATE, ISSUE_ROBOTS, ISSUE_NO_TARGET})


def _is_error_issue(issue: str) -> bool:
    return issue in _ERROR_ISSUES or issue.startswith(ISSUE_ERROR)


def _summarize(results: list[_RedirectRowResult]) -> RedirectSummary:
    correct = sum(1 for row in results if row.is_correct)
    return RedirectSummary(
        total=len(results),
        correct=correct,
        wrong=len(results) - correct,
        chains=sum(1 for row in results if row.chain_length > 1),
        loops=sum(1 for row in results if row.issue == ISSUE_LOOP),
        errors=sum(1 for row in results if _is_error_issue(row.issue)),
    )


# --------------------------------------------------------------------------- #
# Orchestration
# --------------------------------------------------------------------------- #
@dataclass(frozen=True, slots=True)
class _RunContext:
    """The per-run collaborators a worker thread needs, bundled so the job
    function keeps a signature a human can read."""

    options: RedirectCheckOptions
    pool: _SessionPool
    robots: _RobotsCache
    throttle: RedirectThrottle
    pause_flag: threading.Event | None = None
    cancel_flag: threading.Event | None = None

    def stopped(self) -> bool:
        return self.cancel_flag is not None and self.cancel_flag.is_set()


def _wait_while_paused(context: _RunContext) -> None:
    # Waiting HERE, before the row's first request, is what makes Pause mean
    # "stop hitting the site". The old code waited on the consumer side, so
    # every queued worker kept firing requests while the GUI showed "paused".
    # Cancel breaks the wait, or a paused run could never be shut down.
    while context.pause_flag is not None and context.pause_flag.is_set() and not context.stopped():
        time.sleep(0.2)


def _run_redirect_job(job: _RedirectJob, context: _RunContext) -> _RedirectRowResult:
    _wait_while_paused(context)
    with context.throttle.slot():
        session = context.pool.session()
        outcome = _follow_chain(session, job.old_url, context.options, context.robots, context.throttle)
    return _row_result(job, outcome, _verdict(outcome, job.new_url))


def _safe_run(job: _RedirectJob, context: _RunContext) -> _RedirectRowResult:
    """One bad row must never take the run down with it — an unexpected
    exception used to propagate out of the executor, losing every row that had
    already completed."""
    try:
        return _run_redirect_job(job, context)
    except Exception as exc:  # noqa: BLE001 - a row's failure is data, not a crash
        _logger.warning("redirect row failed for %s: %s", job.old_url, exc)
        outcome = RedirectOutcome("", "", (), f"{ISSUE_ERROR}: {exc.__class__.__name__}")
        return _row_result(job, outcome, (False, outcome.issue))


ProgressCallback = Callable[[int, int, str, str | int], None]


def _drain(
    futures: dict[Future[_RedirectRowResult], _RedirectJob],
    df: pd.DataFrame,
    context: _RunContext,
    progress_callback: ProgressCallback | None,
) -> list[_RedirectRowResult]:
    """Apply results in completion order, so progress reflects work actually
    done rather than waiting on the slowest early row."""
    results: list[_RedirectRowResult] = []
    total = len(futures)
    for future in as_completed(futures):
        result = future.result()
        _apply_redirect_result(df, result)
        results.append(result)
        if progress_callback:
            progress_callback(len(results), total, result.old_url[:_MAX_LOG_URL], result.status or result.issue)
        if context.stopped():
            break
    return results


def _execute(
    jobs: list[_RedirectJob],
    df: pd.DataFrame,
    context: _RunContext,
    progress_callback: ProgressCallback | None,
) -> list[_RedirectRowResult]:
    with ThreadPoolExecutor(max_workers=context.throttle.max_parallel) as executor:
        futures = {executor.submit(_safe_run, job, context): job for job in jobs}
        results = _drain(futures, df, context, progress_callback)
        for future in futures:
            future.cancel()
    return results


def check_redirects(
    excel_path: str | Path,
    *,
    options: RedirectCheckOptions | None = None,
    throttle: RedirectThrottle | None = None,
    progress_callback: ProgressCallback | None = None,
    pause_flag: threading.Event | None = None,
    cancel_flag: threading.Event | None = None,
) -> RedirectRunResult:
    """Verify every Old URL -> New URL row and write ``<name>_results.xlsx``.

    Accepts an Excel sheet or a .csv (including the Site Crawl redirect map).
    Rows already finished when a run is cancelled are kept and written, so
    stopping a long audit still leaves a usable report.
    """
    source = Path(excel_path)
    run_options = options or RedirectCheckOptions()
    df = _load_redirect_frame(source)
    _ensure_result_columns(df)
    context = _RunContext(
        options=run_options,
        pool=_SessionPool(run_options.verify_ssl),
        robots=_RobotsCache(run_options.timeout, run_options.verify_ssl),
        throttle=throttle or RedirectThrottle(run_options.max_workers, max_parallel=run_options.max_workers),
        pause_flag=pause_flag,
        cancel_flag=cancel_flag,
    )
    results = _execute(_redirect_jobs(df), df, context, progress_callback)
    _mark_unchecked(df, results)
    target = source.with_name(f"{source.stem}_results.xlsx")
    return RedirectRunResult(_write_redirect_results(df, target), _summarize(results), context.stopped())


def _mark_unchecked(df: pd.DataFrame, results: list[_RedirectRowResult]) -> None:
    """Name the rows a cancelled run never reached. Left blank they read as
    failures — an empty status is exactly what an unreachable host produces —
    so a stopped audit would look like a broken one."""
    if len(results) == len(df):
        return
    checked = {result.index for result in results}
    issue_column = df.columns.get_loc(ISSUE_COLUMN)
    for position in range(len(df)):
        if position not in checked:
            df.iloc[position, issue_column] = ISSUE_NOT_CHECKED
