from __future__ import annotations

from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook

from silentfrog.crawl_run_repository import CrawlRunRef  # type: ignore[reportMissingImports]
from silentfrog.crawl_store import CrawlStore, StoredAudit  # type: ignore[reportMissingImports]
from silentfrog.crawl_types import CrawlPayload  # type: ignore[reportMissingImports]
from silentfrog.exporters import export_site_crawl_report  # type: ignore[reportMissingImports]
from silentfrog.exporters.action_workbook import ACTION_SHEET_NAMES  # type: ignore[reportMissingImports]
from silentfrog.image_diagnostics import IMAGE_HEADERS  # type: ignore[reportMissingImports]
from silentfrog.site_crawl_types import SiteCrawlReport, SiteCrawlResult  # type: ignore[reportMissingImports]


def _payload(
    url: str = "https://example.com/page",
    *,
    redirect_status: str = "200",
    fetch_status: int = 200,
) -> CrawlPayload:
    return CrawlPayload.from_raw(
        {
            "meta": [["title", "Duplicate Title", "15"], ["description", "", "0"]],
            "headers": [["h1", "Duplicate Title"]],
            "images": [
                [
                    "https://example.com/img.png",
                    "Alt",
                    "Title",
                    "image/png",
                    "640",
                    "320",
                    "42 KB",
                    "30m",
                    "Lazy",
                    "High",
                ]
            ],
            "links": [
                [
                    "https://example.com/target",
                    "Anchor",
                    "Interno",
                    "follow",
                    "200",
                    "OK",
                    "Body",
                    "Main Heading",
                    "com",
                ]
            ],
            "schema": {
                "summary": {"total": 1, "by_type": {"Product": 1}, "errors": ["missing price"]},
                "blocks": [
                    {
                        "@type": "Product",
                        "name": "Widget",
                        "_extracted_via": "json-ld",
                        "_schema_errors": ["missing price"],
                    }
                ],
                "eligibility": [
                    {
                        "type": "Product",
                        "detected": True,
                        "count": 1,
                        "eligibility": "Incomplete",
                        "missing_fields": ["offers"],
                        "warnings": ["missing price"],
                    }
                ],
                "issues": [],
            },
            "canonical": {"target": url, "self": True, "multiple": False, "status": "200"},
            "redirect": {"chain": [url], "hops": 0, "final_status": redirect_status, "loop": False},
            "robots": {"*": [["Allow", "/"]]},
            "meta_robots": "index, follow",
            "hreflang": [["en", "https://example.com/en", "200", "Yes", "Yes"]],
            "ai_crawl": [["GPTBot", "GPTBot", "Yes", "-", "-", "Allowed", "Allowed by robots.txt"]],
            "serp": {
                "title": "Duplicate Title",
                "description": "",
                "url": url,
                "site_name": "",
                "breadcrumb": "",
                "favicon": "",
            },
            "serp_audit": {
                "too_long": "No",
                "too_short": "Yes",
                "px_over": "No",
                "px_under": "Yes",
                "equals_h1": "Yes",
                "missing": "No",
                "px_len": "120",
                "char_len": "15",
            },
            "keywords": [
                {
                    "term": "design table",
                    "length": 2,
                    "frequency": 3,
                    "density": 1.5,
                    "in_title": True,
                    "in_description": False,
                    "heading_count": 1,
                    "first_position": 4,
                }
            ],
            "content_quality": {
                "language": "en",
                "word_count": 250,
                "paragraph_count": 4,
                "substantial_paragraph_count": 3,
                "average_words_per_paragraph": 62.5,
                "title_present": True,
                "meta_description_present": False,
                "h1_count": 1,
                "h2_h6_count": 2,
                "title_h1_alignment": "Aligned",
                "intro_paragraph": "Present",
                "thin_content_risk": "Low",
                "heading_structure": "Good",
                "verdict": "Strong",
            },
            "ai_visibility": {
                "summary": {
                    "verdict": "Needs work",
                    "good_count": 1,
                    "warning_count": 1,
                    "critical_count": 0,
                },
                "checks": [
                    {
                        "area": "Access",
                        "check": "AI crawler access",
                        "status": "warning",
                        "details": "Allowed with caveats",
                        "recommendation": "Review robots.txt",
                    }
                ],
            },
            "performance": {
                "status": fetch_status,
                "nav_ttfb_ms": 123,
                "nav_total_ms": 456,
                "transfer_size": 2048,
                "summary": {
                    "verdict": "Needs work",
                    "transfer_size": 2048,
                    "total_resource_bytes": 4096,
                    "total_page_bytes": 6144,
                    "total_resource_count": 2,
                    "third_party_bytes": 0,
                    "third_party_count": 0,
                    "critical_issue_count": 0,
                    "warning_issue_count": 1,
                    "info_issue_count": 0,
                },
                "resource_breakdown": [{"type": "img", "count": 1, "bytes": 4096}],
                "issues": [
                    {
                        "key": "image_weight",
                        "severity": "warning",
                        "message": "Large image",
                        "evidence": "Measured image weight is high.",
                        "recommendation": "Compress image assets.",
                    }
                ],
                "opportunity_details": [{"message": "Use next-gen image formats.", "severity": "warning"}],
                "top_offenders": [
                    {"type": "img", "url": "https://example.com/img.png", "bytes": 4096, "blocking": False}
                ],
                "scripts": {"blocking": {"count": 1, "bytes": 512}, "async": {"count": 2, "bytes": 1024}},
            },
            "social": {
                "open_graph": {
                    "title": "OG title",
                    "description": "OG description",
                    "image": "https://example.com/og.png",
                    "card": "",
                    "image_type": "image/png",
                    "image_width": 1200,
                    "image_height": 630,
                    "issues": [],
                },
                "twitter": {
                    "title": "TW title",
                    "description": "TW description",
                    "image": "https://example.com/tw.png",
                    "card": "summary_large_image",
                    "image_type": "image/png",
                    "image_width": 1200,
                    "image_height": 630,
                    "issues": ["missing alt"],
                },
            },
        }
    )


def test_export_site_crawl_report_creates_summary_and_issue_sheets(tmp_path: Path) -> None:
    first = SiteCrawlResult.from_payload(
        "https://example.com/a", _payload("https://example.com/a", redirect_status="403")
    )
    second = SiteCrawlResult.from_payload("https://example.com/b", _payload("https://example.com/b"))
    failed = SiteCrawlResult.failed("https://example.com/fail", "boom")
    report = SiteCrawlReport.from_results([first, second, failed], discovered_count=3)
    output = tmp_path / "site-crawl.xlsx"

    export_site_crawl_report(report, output)

    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == [
            *ACTION_SHEET_NAMES,
            "Summary",
            "Indexability issues",
            "Meta issues",
            "Structured data",
            "Images",
            "GEO Aggregate",
            "GEO Score",
            "AI Visibility",
            "Errors",
            "Meta detail",
            "Headers detail",
            "Images detail",
            "Links detail",
            "Redirect detail",
            "Canonical detail",
            "Indexability detail",
            "Robots detail",
            "Hreflang detail",
            "Structured detail",
            "Structured eligibility",
            "Content quality detail",
            "Keywords detail",
            "AI crawl detail",
            "GEO Score detail",
            "AI Visibility detail",
            "Performance detail",
            "SERP detail",
            "Social detail",
        ]
        assert workbook["Read me - Legend"]["B4"].value == "Only blockers or high-confidence damage."
        assert workbook["Executive summary"]["A2"].value == "Report type"
        assert workbook["Executive summary"]["B3"].value == 3
        assert workbook["Prioritized issues"]["A1"].value == "Severity"
        assert workbook["Prioritized issues"]["C2"].value == "https://example.com/fail"
        assert workbook["Prioritized issues"]["F2"].value
        assert workbook["Affected URLs"]["A2"].value == "https://example.com/a"
        assert workbook["Content-meta actions"]["D2"].value in {
            "The page has no meta description.",
            "Detected structured data is incomplete for a rich result opportunity.",
        }
        assert workbook["Technical actions"]["D2"].value
        assert workbook["AI-GEO actions"]["D2"].value == "Allowed with caveats"
        assert workbook["Appendix - raw data"]["A2"].value == "Summary"
        assert workbook["Summary"]["A2"].value == "https://example.com/a"
        assert workbook["Summary"]["B2"].value == "200"
        assert workbook["Summary"]["C1"].value == "Redirect status"
        assert workbook["Summary"]["C2"].value == "403"
        assert workbook["Meta issues"]["B2"].value in {"Duplicate title", "Meta description"}
        assert workbook["Structured data"]["C2"].value == "Product"
        assert workbook["Errors"]["A2"].value == "https://example.com/fail"
        assert workbook["Meta detail"]["A2"].value == "https://example.com/a"
        assert workbook["Images detail"]["B1"].value == IMAGE_HEADERS[0]
        assert workbook["Images detail"]["G2"].value == "320"
        assert workbook["Links detail"]["B2"].value == "https://example.com/target"
        assert workbook["Structured detail"]["C2"].value == "Product"
        assert workbook["Structured eligibility"]["E2"].value == "offers"
        assert workbook["AI Visibility detail"]["C2"].value == "AI crawler access"
        assert workbook["Performance detail"]["D2"].value == 2048
        assert workbook["SERP detail"]["C2"].value == "Title"
        assert workbook["Social detail"]["B2"].value == "OpenGraph"
        assert "https://example.com/fail" not in _column_values(workbook["Meta detail"], "A")
    finally:
        workbook.close()


def test_export_site_crawl_report_streams_payload_sheets_from_run_repository(tmp_path: Path) -> None:
    # A store-backed report carries only a CrawlRunRef; the export streams every
    # URL through the run-bound repository and emits the full payload-derived
    # sheets. A store-less inline report with a stripped result has no payload to
    # show, so those sheets stay empty.
    url = "https://example.com/stripped"
    payload = _payload(url)
    db = tmp_path / "crawl.db"
    store = CrawlStore(db)
    run_id = store.start_run("example.com", "https://example.com/", "list")
    store.save_audit(run_id, StoredAudit(url=url, http_status="200", payload=payload.to_mapping()))
    store.finish_run(run_id)
    store.close()

    inline_report = SiteCrawlReport.from_results(
        [replace(SiteCrawlResult.from_payload(url, payload), payload=None)], discovered_count=1
    )
    without = tmp_path / "without.xlsx"
    export_site_crawl_report(inline_report, without)
    workbook = load_workbook(without)
    try:
        assert workbook["Structured data"]["A2"].value == "-"
        assert workbook["GEO Score"]["A2"].value == "-"
    finally:
        workbook.close()

    run_report = SiteCrawlReport.from_run(
        CrawlRunRef(db, run_id),
        discovered_count=1,
        crawled_count=1,
        skipped_count=0,
        failed_count=0,
        base_url="https://example.com/",
    )
    with_repo = tmp_path / "with.xlsx"
    export_site_crawl_report(run_report, with_repo)
    workbook = load_workbook(with_repo)
    try:
        assert workbook["Structured data"]["A2"].value == url
        assert workbook["Structured data"]["C2"].value == "Product"
        assert workbook["GEO Score"]["A2"].value == url
    finally:
        workbook.close()


def _column_values(sheet, column: str) -> list[object]:
    return [cell.value for cell in sheet[column] if cell.value]
